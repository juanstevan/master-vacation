"""Write the extracted homes out as CSV, XLSX, JSONL and a RAG-ready corpus."""
from __future__ import annotations

import csv
import json
import logging
import re
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .records import FIELDS, Record, record_to_row

log = logging.getLogger("mvh.export")

CELL_LIMIT = 32000  # Excel's hard limit is 32767
ILLEGAL = re.compile(r"[\000-\010\013\014\016-\037]")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _safe(value):
    if isinstance(value, str):
        value = ILLEGAL.sub("", value)
        if len(value) > CELL_LIMIT:
            value = value[: CELL_LIMIT - 15] + " ...[truncated]"
    return value


def _style_sheet(ws, widths: Sequence[int], wrap_columns: Sequence[int] = ()) -> None:
    for col_index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_index)].width = width
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    for col_index in wrap_columns:
        letter = get_column_letter(col_index)
        for row in range(2, ws.max_row + 1):
            ws[f"{letter}{row}"].alignment = Alignment(wrap_text=True, vertical="top")


def write_csv(records: list[Record], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig so Excel opens accented characters correctly on double-click.
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, extrasaction="ignore")
        writer.writeheader()
        for record in records:
            writer.writerow(record_to_row(record))
    log.info("wrote %s (%d rows)", path, len(records))


def write_jsonl(records: list[Record], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        for record in records:
            payload = dict(record)
            payload["_amenities"] = record.amenities
            payload["_images"] = record.images
            payload["_fees"] = record.fees
            payload["_extras"] = record.extras
            payload["_sources"] = record.sources
            payload["_notes"] = record.notes
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")
    log.info("wrote %s (%d records)", path, len(records))


def write_xlsx(records: list[Record], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Homes"
    ws.append(FIELDS)
    for record in records:
        ws.append([_safe(v) for v in record_to_row(record).values()])
    widths = []
    for field in FIELDS:
        widths.append(60 if field in ("description", "amenities") else
                      max(12, min(34, len(field) + 6)))
    _style_sheet(ws, widths, wrap_columns=[FIELDS.index("description") + 1,
                                           FIELDS.index("amenities") + 1])

    amen = wb.create_sheet("Amenities")
    amen.append(["listing_id", "house_number", "amenity"])
    for record in records:
        for item in record.amenities:
            amen.append([record.get("listing_id"), record.get("house_number"), _safe(item)])
    _style_sheet(amen, [14, 16, 60])

    fees = wb.create_sheet("Quote details")
    fees.append(["listing_id", "house_number", "label", "value"])
    for record in records:
        for label, value in record.fees:
            fees.append([record.get("listing_id"), record.get("house_number"),
                         _safe(label), _safe(value)])
    _style_sheet(fees, [14, 16, 40, 24])

    extras = wb.create_sheet("Other fields")
    extras.append(["listing_id", "house_number", "field", "value"])
    for record in records:
        for key, value in record.extras.items():
            extras.append([record.get("listing_id"), record.get("house_number"),
                           _safe(key), _safe(value)])
    _style_sheet(extras, [14, 16, 34, 60], wrap_columns=[4])

    prov = wb.create_sheet("Data sources")
    prov.append(["listing_id", "field", "came_from"])
    for record in records:
        for field, source in record.sources.items():
            prov.append([record.get("listing_id"), field, source])
    _style_sheet(prov, [14, 24, 20])

    cover = wb.create_sheet("Coverage")
    cover.append(["field", "filled", "total", "percent_filled"])
    total = len(records) or 1
    for field in FIELDS:
        filled = sum(1 for r in records if str(r.get(field, "")).strip())
        cover.append([field, filled, len(records), round(100.0 * filled / total, 1)])
    _style_sheet(cover, [26, 10, 10, 16])

    wb.save(path)
    log.info("wrote %s (%d homes)", path, len(records))


# ------------------------------------------------------------------ RAG corpus
def _front_matter(record: Record) -> str:
    keys = [
        "house_number", "listing_id", "name", "community", "bedrooms", "bathrooms",
        "sleeps", "pool", "spa", "pets", "check_in_time", "check_out_time", "url",
    ]
    lines = ["---"]
    for key in keys:
        value = record.get(key, "")
        if str(value).strip():
            lines.append(f'{key}: "{str(value).replace(chr(34), chr(39))}"')
    lines.append("---")
    return "\n".join(lines)


def write_rag_corpus(records: list[Record], directory: Path) -> None:
    """One Markdown document per home: what you actually feed the RAG index."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    for record in records:
        slug = str(record.get("house_number") or record.get("listing_id") or "unknown")
        slug = re.sub(r"[^A-Za-z0-9_-]", "_", slug)
        parts = [_front_matter(record), ""]
        title = record.get("name") or f"Home {slug}"
        parts.append(f"# {title}")

        facts = [
            (label, record.get(field))
            for label, field in [
                ("Home number", "house_number"), ("Community", "community"),
                ("Property type", "property_type"), ("Bedrooms", "bedrooms"),
                ("Bathrooms", "bathrooms"), ("Sleeps", "sleeps"),
                ("Square feet", "sqft"), ("Pool", "pool"), ("Spa / hot tub", "spa"),
                ("Pets", "pets"), ("Smoking", "smoking"),
                ("Check-in", "check_in_time"), ("Check-out", "check_out_time"),
                ("Minimum nights", "min_nights"), ("Address", "address"),
            ]
            if str(record.get(field, "")).strip()
        ]
        if facts:
            parts += ["", "## Key facts", ""]
            parts += [f"- **{label}:** {value}" for label, value in facts]

        if str(record.get("description", "")).strip():
            parts += ["", "## Description", "", str(record["description"])]

        if record.amenities:
            parts += ["", "## Amenities", ""]
            parts += [f"- {item}" for item in record.amenities]

        if record.fees:
            parts += ["", "## Quote / fees", ""]
            parts += [f"- **{label}:** {value}" for label, value in record.fees]

        if record.extras:
            parts += ["", "## Additional details", ""]
            parts += [f"- **{key}:** {value}" for key, value in record.extras.items()]

        if str(record.get("quote_text", "")).strip():
            parts += ["", "## Confirmation letter / quote page", "",
                      str(record["quote_text"])]

        if str(record.get("page_text", "")).strip():
            parts += ["", "## Full listing page text", "", str(record["page_text"])]

        parts += ["", "---", f"Source: {record.get('url', '')}",
                  f"Captured: {record.get('scraped_at', '')}"]
        (directory / f"{slug}.md").write_text("\n".join(parts), encoding="utf-8")
    log.info("wrote %d markdown documents to %s", len(records), directory)
