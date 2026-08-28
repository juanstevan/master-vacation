"""End-to-end: serve a fake site, crawl it, check the spreadsheet. python3 tests/test_pipeline.py"""
from __future__ import annotations

import shutil
import sys
import tempfile
import threading
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from http.server import HTTPServer  # noqa: E402

from fixtures.fake_site import Handler  # noqa: E402
from mvh.cli import main  # noqa: E402


def run() -> int:
    server = HTTPServer(("127.0.0.1", 0), Handler)
    port = server.server_address[1]
    threading.Thread(target=server.serve_forever, daemon=True).start()
    workdir = Path(tempfile.mkdtemp(prefix="mvh-e2e-"))

    try:
        code = main([
            "--base-url", f"http://127.0.0.1:{port}",
            "--out", str(workdir), "--delay", "0", "all",
        ])
        assert code == 0, f"pipeline exited {code}"

        import csv
        rows = {
            r["house_number"]: r
            for r in csv.DictReader(open(workdir / "homes.csv", encoding="utf-8-sig"))
        }
        assert set(rows) == {"1906", "2014", "3302"}, sorted(rows)

        # One home per markup style, checked against the fixture's own values.
        assert rows["1906"]["bedrooms"] == "6"          # from JSON-LD
        assert rows["2014"]["bedrooms"] == "4"          # from embedded JS state
        assert rows["3302"]["bedrooms"] == "8"          # from plain HTML divs
        assert rows["1906"]["pool"] == "Private pool"
        assert rows["2014"]["pets"] == "No"
        assert rows["3302"]["community"] == "Example Lakes"
        for row in rows.values():
            assert row["amenity_count"] == "7"
            assert row["check_in_time"] == "4:00 PM"
            assert row["quote_total"] == "$1,713.40"
        # Columns the real site never publishes were dropped from the workbook.
        assert "confirmation_number" not in rows["1906"], sorted(rows["1906"])
        assert "rating" not in rows["1906"]

        for name in ("homes.xlsx", "homes.csv", "parsed.jsonl"):
            assert (workdir / name).exists(), f"missing {name}"
        docs = sorted(p.name for p in (workdir / "rag_corpus").glob("*.md"))
        assert docs == ["1906.md", "2014.md", "3302.md"], docs

        from openpyxl import load_workbook
        wb = load_workbook(workdir / "homes.xlsx")
        assert wb.sheetnames == ["Homes", "Amenities", "Quote details",
                                 "Other fields", "Data sources", "Coverage"]
        assert wb["Homes"].max_row == 4          # header + 3 homes
        assert wb["Amenities"].max_row == 22     # header + 3 x 7 amenities

        print("  PASS  end-to-end: discover -> fetch -> parse -> export")
        print("\n1/1 passed")
        return 0
    except AssertionError as exc:
        print(f"  FAIL  end-to-end: {exc}")
        return 1
    finally:
        server.shutdown()
        shutil.rmtree(workdir, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(run())
